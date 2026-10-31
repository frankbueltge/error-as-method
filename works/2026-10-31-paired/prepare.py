"""Session 165 (experiment 45, Paired): keys, contact sets, picture sheets and briefs.

Reads the AAPS release (Nittono & Nakamura 2026, https://osf.io/3uzrw, CC BY 4.0) from a download directory that holds
AAPS_dataset.xlsx and AAPS_images.zip, checks their SHA-256 against sources/MANIFEST.json, and writes:
  people.json      per key: file, species, source URL, n, the people's kawaii mean, median and SD (derived values)
  plan.json        the 72 keys, which 24 are contact keys, the shuffled pairing, each maker's arrangement
  sheets/*.jpg     4 x 3 picture sheets, numbered tiles (scored: three arrangements; contact: two)
  thumbs/*.jpg     240 px pictures for the face (all in the public domain or CC0, per the release)
  briefs/*.txt     the 24 briefs
Run: python3 prepare.py <download dir>
"""
import hashlib, io, json, random, statistics as st, sys, zipfile
from pathlib import Path
import openpyxl
from PIL import Image, ImageDraw, ImageFont

SEED = 165
HERE = Path(__file__).parent
DL = Path(sys.argv[1])
MAN = json.loads((HERE / "sources" / "MANIFEST.json").read_text())
for f in MAN["files"]:
    h = hashlib.sha256((DL / f["name"]).read_bytes()).hexdigest()
    assert h == f["sha256"], (f["name"], h)

wb = openpyxl.load_workbook(DL / "AAPS_dataset.xlsx", read_only=True)
rows = list(wb["800 rated"].iter_rows(values_only=True))
head = rows[0]
ix = {k: head.index(k) for k in ("Filename", "Rating n", "M_Kawaii", "SD_Kawaii", "Species", "URL")}
meta = {}
for r in rows[1:]:
    if r[ix["Filename"]] and isinstance(r[ix["M_Kawaii"]], (int, float)):
        meta[r[ix["Filename"]]] = dict(file=r[ix["Filename"]], species=r[ix["Species"]], url=r[ix["URL"]],
                                       n=r[ix["Rating n"]], mean=r[ix["M_Kawaii"]], sd=r[ix["SD_Kawaii"]])
raw = list(wb["N=914"].iter_rows(values_only=True))
cols = [(j, raw[1][j]) for j, h in enumerate(raw[0]) if isinstance(h, str) and h.endswith("_Kawaii")]
vals = {f: [] for _, f in cols}
for r in raw[2:]:
    for j, f in cols:
        if r[j] is not None:
            vals[f].append(float(r[j]))
for f, m in meta.items():
    v = vals[f]
    assert len(v) == m["n"] and abs(st.mean(v) - m["mean"]) < 1e-9, f
    m["median"] = st.median(v)
assert len(meta) == 800

rng = random.Random(SEED)
keys = rng.sample(sorted(meta), 72)
contact, scored = keys[:24], keys[24:]
kid = {f: f"k{i+1:02d}" for i, f in enumerate(keys)}


def spearman(a, b):
    def rank(x):
        o = sorted(range(len(x)), key=lambda i: x[i]); r = [0.0] * len(x); i = 0
        while i < len(o):
            j = i
            while j + 1 < len(o) and x[o[j + 1]] == x[o[i]]: j += 1
            for t in range(i, j + 1): r[o[t]] = (i + j) / 2
            i = j + 1
        return r
    ra, rb = rank(a), rank(b); ma, mb = st.mean(ra), st.mean(rb)
    num = sum((x - ma) * (y - mb) for x, y in zip(ra, rb))
    return num / (sum((x - ma) ** 2 for x in ra) * sum((y - mb) ** 2 for y in rb)) ** 0.5


# The shuffled pairing for arm S: a derangement of the contact keys' (median, mean) pairs, the first one from seed 165
# upward whose Spearman rho with the true means lies within +-0.1. Every number told to S is true of some picture.
s = SEED
while True:
    r2 = random.Random(s); perm = contact[:]; r2.shuffle(perm)
    if all(a != b for a, b in zip(contact, perm)):
        rho = spearman([meta[f]["mean"] for f in contact], [meta[f]["mean"] for f in perm])
        if abs(rho) <= 0.1: break
    s += 1
shuffle_seed, shuffle_rho = s, rho
told = {"B": {f: f for f in contact}, "S": dict(zip(contact, perm))}

# Arrangements: scored keys in three seeded orders, contact keys in two. Maker i of an arm uses (i-1) % 3 and (i-1) % 2.
arr_s = [random.Random(SEED * 10 + a).sample(scored, 48) for a in range(3)]
arr_c = [random.Random(SEED * 100 + a).sample(contact, 24) for a in range(2)]

z = zipfile.ZipFile(DL / "AAPS_images.zip")
font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 26)
TILE, PAD, LAB = 300, 12, 38


def sheet(files, labels, out):
    W, H = 4 * TILE + 5 * PAD, 3 * (TILE + LAB) + 4 * PAD
    im = Image.new("RGB", (W, H), "white"); d = ImageDraw.Draw(im)
    for i, (f, lab) in enumerate(zip(files, labels)):
        x, y = PAD + (i % 4) * (TILE + PAD), PAD + (i // 4) * (TILE + LAB + PAD)
        d.text((x, y + 2), lab, fill="black", font=font)
        p = Image.open(io.BytesIO(z.read(f))).convert("RGB").resize((TILE, TILE), Image.LANCZOS)
        im.paste(p, (x, y + LAB))
    im.save(out, quality=88)


(HERE / "sheets").mkdir(exist_ok=True); (HERE / "thumbs").mkdir(exist_ok=True)
for a, order in enumerate(arr_s):
    for s_ in range(4):
        sheet(order[s_ * 12:(s_ + 1) * 12], [str(s_ * 12 + i + 1) for i in range(12)], HERE / "sheets" / f"scored-{a+1}-{s_+1}.jpg")
for a, order in enumerate(arr_c):
    for s_ in range(2):
        sheet(order[s_ * 12:(s_ + 1) * 12], [f"T{s_ * 12 + i + 1}" for i in range(12)], HERE / "sheets" / f"contact-{a+1}-{s_+1}.jpg")
for f in keys:
    Image.open(io.BytesIO(z.read(f))).convert("RGB").resize((240, 240), Image.LANCZOS).save(HERE / "thumbs" / f"{kid[f]}.jpg", quality=82)

SHEETDIR = "/home/user/error-as-method/works/2026-10-31-paired/sheets"
PEOPLE = ("In an online study in Japan in March 2026, 914 adults whose first language is Japanese (recruited through a Japanese "
          "crowdsourcing service, average age 46) each saw 50 pictures of this kind, one at a time for 3 seconds, and then rated each "
          "one on several scales. One scale read, in Japanese, 「この写真は・・・」 (\"This image is ...\"), from 1 = "
          "全くかわいくない (not kawaii at all) to 9 = 非常にかわいい (extremely kawaii), with a whole number. "
          "Each picture was rated by about 57 of them (35 to 85).")
LOOK = ("Look at the pictures with your file-reading tool. They are on four sheets, twelve numbered pictures each:\n"
        "{s}\n\nUse no other tool, program, search or lookup of any kind, and read no file other than the picture sheets named in this message. Answer from what you see and what is written here.")
ANSWER_A = ("For each of the 48 numbered pictures, rate it yourself: how kawaii is this image to you, from 1 (not kawaii at all) "
            "to 9 (extremely kawaii), with a whole number.")
ANSWER_D = ("For each of the 48 numbered pictures, do not rate it yourself. State how those people rated it, on their own 1 to 9 "
            "scale: give the median of their ratings and their mean rating (decimals allowed).")
FORM_A = "Hand back exactly 48 lines in the form\n1: <rating>\n2: <rating>\n... and then one more line:\nHOW: <one sentence on how you rated>"
FORM_D = ("Hand back exactly 48 lines in the form\n1: median <number>, mean <number>\n2: median <number>, mean <number>\n... and then two more lines:\n"
          "HOW: <one sentence on how you arrived at the numbers>\nKNOWN: <yes, no or partly> - did you know this study or its data before today?")


def contact_block(arm, a):
    order = arr_c[a]; tm = told[arm]
    lines = [f"T{i+1}: median {fmt(meta[tm[f]]['median'])}, mean {meta[tm[f]]['mean']:.2f}" for i, f in enumerate(order)]
    paths = "\n".join(f"{SHEETDIR}/contact-{a+1}-{s_+1}.jpg" for s_ in range(2))
    return ("Before that, here are 24 other pictures from the same study, on two sheets labelled T1 to T24:\n" + paths +
            "\nFor each of them, these are the median and the mean of the people's ratings on the 1 to 9 scale:\n" + "\n".join(lines) + "\n\n")


def fmt(x):
    return str(int(x)) if x == int(x) else f"{x:.1f}"


plan = {"seed": SEED, "keys": [{"id": kid[f], "file": f, "role": "contact" if f in contact else "scored"} for f in keys],
        "shuffle": {"seed": shuffle_seed, "rho_with_true_means": round(shuffle_rho, 4),
                    "told_S": {kid[f]: kid[g] for f, g in told["S"].items()}},
        "arrangements_scored": [[kid[f] for f in o] for o in arr_s],
        "arrangements_contact": [[kid[f] for f in o] for o in arr_c], "makers": {}}
(HERE / "briefs").mkdir(exist_ok=True)
for arm in "ADBS":
    for i in range(1, 7):
        a, c = (i - 1) % 3, (i - 1) % 2
        paths = "\n".join(f"{SHEETDIR}/scored-{a+1}-{s_+1}.jpg" for s_ in range(4))
        if arm == "A":
            text = ("Below are pictures of animals.\n\n" + LOOK.format(s=paths) + "\n\n" + ANSWER_A + "\n\n" + FORM_A)
        else:
            text = ("Below are pictures of animals.\n\n" + PEOPLE + "\n\n" + LOOK.format(s=paths) + "\n\n" +
                    (contact_block(arm, c) if arm in "BS" else "") + ANSWER_D + "\n\n" + FORM_D)
        (HERE / "briefs" / f"{arm}{i}.txt").write_text(text + "\n")
        plan["makers"][f"{arm}{i}"] = {"scored_arrangement": a + 1, "contact_arrangement": (c + 1) if arm in "BS" else None}
(HERE / "plan.json").write_text(json.dumps(plan, indent=1))
(HERE / "people.json").write_text(json.dumps({kid[f]: {**meta[f], "role": "contact" if f in contact else "scored"} for f in keys},
                                             indent=1, ensure_ascii=False))
print("shuffle seed", shuffle_seed, "rho", round(shuffle_rho, 3))
sm = [meta[f]["mean"] for f in scored]; cm = [meta[f]["mean"] for f in contact]
print("scored means", round(min(sm), 2), round(max(sm), 2), "SD", round(st.pstdev(sm), 2), "| contact SD", round(st.pstdev(cm), 2))
print("species scored:", sorted(set(meta[f]["species"] for f in scored))[:60])
